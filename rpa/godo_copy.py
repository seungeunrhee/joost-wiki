# -*- coding: utf-8 -*-
"""
유스트 고도몰 상품 복사등록 도우미
- 기존 상품을 고도몰 [선택 복사] 기능으로 복제 (이미지/상세페이지/옵션까지 서버가 복제)
- 복사본의 상품코드 / 상품명 / 정가 / 판매가 / 재고만 엑셀 값으로 덮어씀
- 썸네일과 상세페이지 수정, 최종 [저장]은 사람이
"""
import os, sys, threading, traceback
from urllib.parse import quote
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP = "유스트 고도몰 상품 복사등록 도우미"

# 엑셀 헤더 -> 내부 키
ALIAS = {
    "원본상품번호": "srcNo", "원본번호": "srcNo", "원본goodsno": "srcNo",
    "원본상품코드": "srcCd", "원본자체상품코드": "srcCd",
    "원본상품명": "srcNm",
    "새상품코드": "goodsCd", "상품코드": "goodsCd", "자체상품코드": "goodsCd",
    "새상품명": "goodsNm", "상품명": "goodsNm",
    "정가": "fixedPrice",
    "판매가": "goodsPrice", "가격": "goodsPrice",
    "재고": "stockCnt", "재고수량": "stockCnt",
}
TEMPLATE_COLS = ["원본상품번호", "원본상품명", "새상품코드", "새상품명", "정가", "판매가", "재고"]

LOGGED_IN_JS = "() => !!document.getElementById('headerSearchKeyword')"

# 복사본 수정 페이지에서 덮어쓸 값
FILL_JS = """
(it) => {
  const ok=[], ng=[];
  const setV=(name,val,label)=>{
    if(val===null||val===undefined||val==='') return;
    const e=document.querySelector("[name='"+name+"']");
    if(!e){ ng.push(label); return; }
    e.focus(); e.value=String(val);
    e.dispatchEvent(new Event('input',{bubbles:true}));
    e.dispatchEvent(new Event('change',{bubbles:true}));
    e.blur(); ok.push(label);
  };
  setV('goodsCd', it.goodsCd, '상품코드');
  setV('goodsNm', it.goodsNm, '상품명');
  setV('fixedPrice', it.fixedPrice, '정가');
  setV('goodsPrice', it.goodsPrice, '판매가');
  if(it.stockCnt!==null && it.stockCnt!==undefined && it.stockCnt!==''){
    const sc=document.querySelector("[name='stockCnt']");
    if(!sc || sc.disabled){
      ng.push('재고 → 옵션 상품입니다. 옵션표에 직접 입력해 주세요');
    } else {
      const r=document.querySelector("input[name='stockFl'][value='y']");
      if(r && !r.checked){ r.click(); }
      setV('stockCnt', it.stockCnt, '재고');
    }
  }
  return {ok:ok, ng:ng};
}
"""

GET_NOS_JS = """() => [...document.querySelectorAll("input[name^='goodsNo[']")].map(e=>parseInt(e.value,10)).filter(n=>!isNaN(n))"""


def norm(s):
    return str(s or "").replace(" ", "").replace("(", "").replace(")", "").lower()


def read_excel(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], "빈 파일입니다."
    head = rows[0]
    idx = {}
    for i, h in enumerate(head):
        k = ALIAS.get(norm(h))
        if k and k not in idx:
            idx[k] = i
    if "srcNo" not in idx and "srcCd" not in idx and "srcNm" not in idx:
        return [], "원본상품번호 / 원본상품코드 / 원본상품명 중 하나는 있어야 합니다."
    items = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        it = {}
        for k, i in idx.items():
            v = r[i] if i < len(r) else None
            if v is None or str(v).strip() == "":
                continue
            if k in ("fixedPrice", "goodsPrice", "stockCnt", "srcNo"):
                try:
                    v = int(float(str(v).replace(",", "").strip()))
                except Exception:
                    v = str(v).strip()
            else:
                v = str(v).strip()
            it[k] = v
        if it:
            items.append(it)
    return items, None


def make_template(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "상품복사"
    ws.append(TEMPLATE_COLS)
    ws.append([665, "", "JUST-HC-030", "허브 크림 30ml", 36000, 32400, 100])
    ws.append(["", "허브 크림 스타트 세트", "JUST-SET-07", "", 108000, 103200, 50])
    for i, w in enumerate([14, 30, 18, 30, 12, 12, 8], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path)


class App:
    def __init__(self, root):
        self.root = root
        root.title(APP)
        root.geometry("760x560")
        self.items = []
        self.stop = False
        self.next_evt = threading.Event()

        top = tk.Frame(root, padx=12, pady=10)
        top.pack(fill="x")
        tk.Label(top, text="고도몰 관리자 주소").grid(row=0, column=0, sticky="w")
        self.url = tk.Entry(top, width=52)
        self.url.insert(0, "https://gdadmin-justartr0736.godomall.com")
        self.url.grid(row=0, column=1, columnspan=3, sticky="w", padx=6)

        tk.Button(top, text="엑셀 양식 만들기", command=self.tpl).grid(row=1, column=0, pady=8, sticky="w")
        tk.Button(top, text="엑셀 열기", width=12, command=self.pick).grid(row=1, column=1, pady=8, sticky="w", padx=6)
        self.fname = tk.Label(top, text="선택된 파일 없음", fg="#666")
        self.fname.grid(row=1, column=2, columnspan=2, sticky="w")

        bar = tk.Frame(root, padx=12)
        bar.pack(fill="x")
        self.start = tk.Button(bar, text="▶ 시작", width=12, height=2, command=self.go, state="disabled")
        self.start.pack(side="left")
        self.nbtn = tk.Button(bar, text="저장했어요 · 다음 ▶", width=20, height=2,
                              command=lambda: self.next_evt.set(), state="disabled")
        self.nbtn.pack(side="left", padx=8)
        tk.Button(bar, text="■ 중지", width=8, height=2, command=self.halt).pack(side="left")
        self.stat = tk.Label(bar, text="대기 중", fg="#0a6", font=("맑은 고딕", 10, "bold"))
        self.stat.pack(side="right")

        self.log = tk.Text(root, height=22, bg="#111", fg="#ddd", font=("Consolas", 9))
        self.log.pack(fill="both", expand=True, padx=12, pady=10)
        self.say("1) 엑셀 양식 만들기 → 값 채우기 → 2) 엑셀 열기 → 3) 시작")
        self.say("   시작하면 크롬이 열립니다. 고도몰 관리자에 직접 로그인해 주세요.")
        self.say("   (아이디·비밀번호는 이 프로그램에 저장되지 않습니다)")

    def say(self, s):
        self.log.insert("end", s + "\n")
        self.log.see("end")
        self.root.update_idletasks()

    def tpl(self):
        p = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                         initialfile="상품복사_양식.xlsx",
                                         filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            make_template(p)
            messagebox.showinfo(APP, "양식을 만들었습니다.\n\n" + p)
        except Exception as e:
            messagebox.showerror(APP, str(e))

    def pick(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not p:
            return
        items, err = read_excel(p)
        if err:
            messagebox.showerror(APP, err)
            return
        if not items:
            messagebox.showerror(APP, "읽을 행이 없습니다.")
            return
        self.items = items
        self.fname.config(text=os.path.basename(p) + "  (%d건)" % len(items), fg="#000")
        self.start.config(state="normal")
        self.say("\n엑셀 %d건 읽음" % len(items))
        for i, it in enumerate(items, 1):
            src = it.get("srcNo") or it.get("srcCd") or it.get("srcNm")
            self.say("  %d. 원본[%s] → %s / %s" % (i, src, it.get("goodsCd", "-"), it.get("goodsNm", "(이름 그대로)")))

    def halt(self):
        self.stop = True
        self.next_evt.set()
        self.say("\n중지 요청됨")

    def go(self):
        self.start.config(state="disabled")
        self.stop = False
        threading.Thread(target=self.work, daemon=True).start()

    # ---------- 작업 ----------
    def work(self):
        try:
            self._work()
        except Exception:
            self.say("\n[오류]\n" + traceback.format_exc())
            self.stat.config(text="오류", fg="#c00")
        finally:
            self.nbtn.config(state="disabled")
            self.start.config(state="normal")

    def _work(self):
        from playwright.sync_api import sync_playwright
        base = self.url.get().strip().rstrip("/")
        goods = base + "/goods"

        def lurl(key, kw, n=100):
            return "%s/goods_list.php?searchFl=y&key=%s&keyword=%s&pageNum=%d&sort=%s" % (
                goods, key, quote(str(kw)), n, quote("g.goodsNo desc"))

        with sync_playwright() as p:
            br = p.chromium.launch(headless=False, channel="chrome", args=["--start-maximized"])
            ctx = br.new_context(no_viewport=True)
            pg = ctx.new_page()
            pg.goto(base)

            self.stat.config(text="로그인 대기", fg="#c60")
            self.say("\n크롬에서 로그인해 주세요. (최대 10분 대기)")
            for _ in range(600):
                if self.stop:
                    return
                try:
                    if pg.evaluate(LOGGED_IN_JS):
                        break
                except Exception:
                    pass
                pg.wait_for_timeout(1000)
            else:
                self.say("로그인이 확인되지 않아 중단합니다.")
                return
            self.say("로그인 확인됨.")
            self.stat.config(text="진행 중", fg="#06c")

            poll = ctx.new_page()

            for i, it in enumerate(self.items, 1):
                if self.stop:
                    break
                label = it.get("goodsCd") or it.get("goodsNm") or "-"
                self.say("\n[%d/%d] %s" % (i, len(self.items), label))

                src = self.find_src(pg, lurl, it)
                if not src:
                    continue
                self.say("   원본 상품번호 %s" % src)

                before = max(poll_nos(poll, lurl) or [0])

                # 복사 실행
                pg.goto(lurl("goodsNo", src), wait_until="domcontentloaded")
                pg.wait_for_timeout(800)
                try:
                    pg.check("input[name='goodsNo[%s]']" % src)
                except Exception:
                    self.say("   ! 목록에서 원본을 찾지 못했습니다. 건너뜁니다.")
                    continue
                pg.click("button.js-check-copy")
                pg.wait_for_selector("#layerGoodsSaleFrm", timeout=20000)
                pg.wait_for_timeout(500)
                # 노출안함 · 판매안함 으로 복사 (작업 끝난 뒤 직접 켜세요)
                pg.check("#goodsDisplay")
                pg.check("#goodsSell")
                pg.wait_for_timeout(300)
                for nm in ("goodsDisplayFl", "goodsDisplayMobileFl", "goodsSellFl", "goodsSellMobileFl"):
                    try:
                        pg.check("#layerGoodsSaleFrm input[name='%s'][value='n']" % nm)
                    except Exception:
                        pass
                pg.click("#goods_saleLayer input.goods-sale-btn")
                pg.wait_for_timeout(700)
                try:
                    pg.click("div.bootstrap-dialog button:has-text('확인')", timeout=10000)
                except Exception:
                    self.say("   ! 복사 확인창을 찾지 못했습니다. 건너뜁니다.")
                    continue
                self.say("   복사 중... (이미지까지 복제되어 시간이 걸립니다)")

                new = None
                for _ in range(80):   # 최대 4분
                    if self.stop:
                        break
                    poll.wait_for_timeout(3000)
                    nos = poll_nos(poll, lurl)
                    if nos and max(nos) > before:
                        new = max(nos)
                        break
                if not new:
                    self.say("   ! 복사본을 찾지 못했습니다. 고도몰 화면을 확인해 주세요.")
                    continue
                self.say("   복사본 생성됨 → 상품번호 %d" % new)

                pg.goto("%s/goods_register.php?goodsNo=%d" % (goods, new), wait_until="domcontentloaded")
                pg.wait_for_timeout(1500)
                r = pg.evaluate(FILL_JS, it)
                self.say("   채움: %s" % (", ".join(r["ok"]) or "없음"))
                if r["ng"]:
                    self.say("   못 채움: %s" % ", ".join(r["ng"]))

                self.say("   → 썸네일·상세페이지를 수정하고 [저장]을 누른 뒤,")
                self.say("     아래 [저장했어요 · 다음 ▶] 을 눌러주세요.")
                self.stat.config(text="사람 차례", fg="#c60")
                self.nbtn.config(state="normal")
                self.next_evt.clear()
                self.next_evt.wait()
                self.nbtn.config(state="disabled")
                self.stat.config(text="진행 중", fg="#06c")

            self.say("\n끝났습니다. 브라우저는 그대로 둡니다.")
            self.say("복사본은 '노출안함·판매안함' 으로 만들어집니다. 확인 후 직접 켜주세요.")
            self.stat.config(text="완료", fg="#0a6")

    def find_src(self, pg, lurl, it):
        if it.get("srcNo"):
            return str(it["srcNo"])
        key, kw = ("goodsCd", it.get("srcCd")) if it.get("srcCd") else ("goodsNm", it.get("srcNm"))
        if not kw:
            self.say("   ! 원본 정보가 없습니다. 건너뜁니다.")
            return None
        pg.goto(lurl(key, kw), wait_until="domcontentloaded")
        pg.wait_for_timeout(800)
        rows = pg.evaluate("""() => [...document.querySelectorAll("input[name^='goodsNo[']")].map(cb=>{
            const tr=cb.closest('tr');
            return {no:cb.value, nm:(tr?tr.innerText:'').replace(/\\s+/g,' ').slice(0,60)};
        })""")
        if not rows:
            self.say("   ! '%s' 로 검색된 상품이 없습니다. 건너뜁니다." % kw)
            return None
        if len(rows) > 1:
            self.say("   ! '%s' 로 %d건이 검색되어 어느 것인지 알 수 없습니다." % (kw, len(rows)))
            for r in rows[:8]:
                self.say("       %s  %s" % (r["no"], r["nm"]))
            self.say("     엑셀의 [원본상품번호] 칸에 위 번호를 적어주세요.")
            return None
        return rows[0]["no"]


def poll_nos(poll, lurl):
    try:
        poll.goto(lurl("goodsNm", ""), wait_until="domcontentloaded")
        poll.wait_for_timeout(400)
        return poll.evaluate(GET_NOS_JS)
    except Exception:
        return []


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
