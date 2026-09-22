# -*- coding: utf-8 -*-
"""
유스트 고도몰 상품등록 자동화
- 엑셀에 적은 상품을 고도몰 관리자 상품등록 화면에 자동으로 채웁니다.
- 로그인은 열린 브라우저에서 직접 하십니다. (비밀번호는 이 프로그램을 거치지 않습니다)
- 마지막 저장 버튼은 사람이 누릅니다.
"""
import os, sys, json, threading, traceback
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP = "유스트 고도몰 상품등록 자동화"
VER = "1.0.0"
CFG = os.path.join(os.path.expanduser("~"), ".just_godo_rpa.json")

COLS = [
    ("상품명", "name"), ("상품코드", "code"), ("카테고리코드", "cate"),
    ("검색키워드", "keyword"), ("노출", "display"), ("판매", "sell"),
    ("재고관리", "useStock"), ("재고수량", "stock"),
    ("정가", "fixedPrice"), ("판매가", "price"), ("매입가", "costPrice"), ("공급가", "supplyPrice"),
    ("브랜드", "brand"), ("제조사", "maker"), ("원산지", "origin"), ("모델명", "model"),
    ("면세", "taxFree"), ("네이버노출", "naver"),
    ("간단설명", "short"), ("이벤트문구", "event"), ("관리자메모", "memo"),
    ("PC상세설명", "descPc"), ("모바일상세설명", "descMobile"),
]
BOOLS = {"display", "sell", "useStock", "taxFree", "naver"}
YES = {"o", "O", "y", "Y", "예", "1", "true", "TRUE", "True", "ㅇ", "사용", "함"}

FILL_JS = r"""
(it) => {
  const fire=(el,ts)=>{(ts||['input','change','blur']).forEach(t=>el.dispatchEvent(new Event(t,{bubbles:true})));
    try{if(window.jQuery)jQuery(el).trigger('change');}catch(e){}};
  const byName=n=>document.getElementsByName(n);
  const sT=(n,v)=>{if(v===undefined||v===null||v==='')return false;const e=byName(n)[0];if(!e)return false;e.value=String(v);fire(e);return true;};
  const sR=(n,v)=>{if(v===undefined||v===null)return false;const l=byName(n);let h=false;
    for(const e of l){if(e.value===String(v)){if(!e.checked){e.checked=true;e.click();}fire(e,['change']);h=true;}}return h;};
  const sS=(i,v)=>{if(v===undefined||v===null||v==='')return false;const e=document.getElementById(i)||byName(i)[0];if(!e)return false;e.value=String(v);fire(e,['change']);return true;};
  const sE=(id,h)=>{if(!h)return false;try{const ta=document.getElementById(id);if(ta)ta.value=h;
    if(window.oEditors&&oEditors.getById&&oEditors.getById[id]){oEditors.getById[id].exec('SET_IE_ALL_CONTENT',[h]);return true;}return !!ta;}catch(e){return false;}};
  const ok=[],ng=[];const m=(r,l)=>{(r?ok:ng).push(l);};
  m(sS('cateGoods1',it.cate),'카테고리');
  m(sT('goodsNm',it.name),'상품명');
  m(sT('goodsCd',it.code),'상품코드');
  m(sT('goodsSearchWord',it.keyword),'검색키워드');
  if(it.display!==undefined) m(sR('goodsDisplayFl',it.display?'y':'n'),'노출상태');
  if(it.sell!==undefined) m(sR('goodsSellFl',it.sell?'y':'n'),'판매상태');
  if(it.useStock!==undefined||it.stock!==undefined){
    const us=(it.useStock===undefined)?true:!!it.useStock;
    m(sR('stockFl',us?'y':'n'),'재고방식');
    if(us){ m(true,'재고수량'); }
  }
  m(sT('fixedPrice',it.fixedPrice),'정가');
  m(sT('goodsPrice',it.price),'판매가');
  m(sT('costPrice',it.costPrice),'매입가');
  m(sT('supplyPrice',it.supplyPrice),'공급가');
  m(sT('brandCdNm',it.brand),'브랜드');
  m(sT('makerNm',it.maker),'제조사');
  m(sT('originNm',it.origin),'원산지');
  m(sT('goodsModelNo',it.model),'모델명');
  if(it.taxFree!==undefined) m(sR('taxFreeFl',it.taxFree?'f':'t'),'과세면세');
  if(it.naver!==undefined) m(sR('naverFl',it.naver?'y':'n'),'네이버노출');
  m(sT('shortDescription',it.short),'간단설명');
  m(sT('eventDescription',it.event),'이벤트문구');
  m(sT('memo',it.memo),'관리자메모');
  m(sE('editor',it.descPc),'PC상세설명');
  m(sE('editor2',it.descMobile||it.descPc),'모바일상세설명');
  return {ok:ok.filter(Boolean),ng:ng};
}
"""

LOGGED_IN_JS = "() => !!document.getElementById('headerSearchKeyword')"


def load_cfg():
    try:
        with open(CFG, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_cfg(d):
    try:
        with open(CFG, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False)
    except Exception:
        pass


def truthy(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip()
    if s == "":
        return None
    return s in YES


def read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [(str(c).strip() if c is not None else "") for c in rows[0]]
    idx = {}
    for ko, key in COLS:
        if ko in head:
            idx[key] = head.index(ko)
    items = []
    for r in rows[1:]:
        if not any(c not in (None, "") for c in r):
            continue
        it = {}
        for key, i in idx.items():
            v = r[i] if i < len(r) else None
            if v is None or str(v).strip() == "":
                continue
            if key in BOOLS:
                b = truthy(v)
                if b is not None:
                    it[key] = b
            elif key in ("stock", "fixedPrice", "price", "costPrice", "supplyPrice"):
                try:
                    it[key] = int(float(str(v).replace(",", "").strip()))
                except Exception:
                    it[key] = str(v).strip()
            else:
                it[key] = str(v).strip()
        if it.get("name"):
            items.append(it)
    return items


def write_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "상품목록"
    ws.append([c[0] for c in COLS])
    fill = PatternFill("solid", fgColor="00519E")
    for i, _ in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(color="FFFFFF", bold=True, size=10)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = 16
    ws.append(["스위스 유스트 허브 크림 30ml", "JK-HC-030", "025", "허브크림,핸드크림",
               "O", "O", "O", 100, 42000, 38000, 21000, "",
               "스위스 유스트", "Just Schweiz AG", "스위스", "",
               "X", "O", "알프스 허브로 만든 데일리 크림", "", "",
               "<p>상세설명</p>", ""])
    ws2 = wb.create_sheet("카테고리코드")
    ws2.append(["코드", "이름"])
    for code, nm in [("022", "프로모션"), ("008", "페이스 케어"), ("025", "허브 크림"),
                     ("026", "허브 에센셜 오일"), ("019", "BATH"), ("024", "바스&스파"),
                     ("023", "SHOWER"), ("020", "바디 케어"), ("016", "헤어 케어"),
                     ("018", "리빙"), ("017", "ALL")]:
        ws2.append([code, nm])
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 22
    ws3 = wb.create_sheet("작성안내")
    for line in ["O / X 로 적는 칸 : 노출, 판매, 재고관리, 면세, 네이버노출",
                 "재고관리 O = 재고량에 따름, X = 무한정 판매",
                 "면세 O = 면세상품, X = 과세상품",
                 "카테고리코드 : 옆 시트 참고 (숫자 3자리, 앞의 0도 그대로)",
                 "상품코드 : 고도몰 코드. 다른 채널 판매자코드에도 같은 값을 씁니다",
                 "빈 칸으로 두면 그 항목은 건너뜁니다",
                 "모바일상세설명을 비우면 PC상세설명과 같은 내용이 들어갑니다"]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 70
    wb.save(path)


class App:
    def __init__(self, root):
        self.root = root
        self.cfg = load_cfg()
        self.items = []
        self.idx = 0
        self.stop = False
        self.next_evt = threading.Event()
        root.title("%s  v%s" % (APP, VER))
        root.geometry("720x560")
        root.configure(bg="#f2f6fb")

        top = tk.Frame(root, bg="#00519E", height=54)
        top.pack(fill="x")
        tk.Label(top, text="고도몰 상품등록 자동화", bg="#00519E", fg="#fff",
                 font=("맑은 고딕", 13, "bold")).pack(side="left", padx=16, pady=13)

        f1 = tk.Frame(root, bg="#f2f6fb")
        f1.pack(fill="x", padx=16, pady=(14, 6))
        tk.Label(f1, text="관리자 주소", bg="#f2f6fb", width=10, anchor="w").pack(side="left")
        self.url = tk.Entry(f1)
        self.url.insert(0, self.cfg.get("url", "https://gdadmin-justartr0736.godomall.com"))
        self.url.pack(side="left", fill="x", expand=True)

        f2 = tk.Frame(root, bg="#f2f6fb")
        f2.pack(fill="x", padx=16, pady=6)
        tk.Label(f2, text="엑셀 파일", bg="#f2f6fb", width=10, anchor="w").pack(side="left")
        self.xl = tk.Entry(f2)
        self.xl.pack(side="left", fill="x", expand=True)
        tk.Button(f2, text="열기", width=7, command=self.pick).pack(side="left", padx=(6, 0))
        tk.Button(f2, text="양식 만들기", width=11, command=self.make_tpl).pack(side="left", padx=(6, 0))

        f3 = tk.Frame(root, bg="#f2f6fb")
        f3.pack(fill="x", padx=16, pady=6)
        self.slow = tk.BooleanVar(value=True)
        tk.Checkbutton(f3, text="천천히 보여주기 (입력되는 과정을 눈으로 확인)",
                       variable=self.slow, bg="#f2f6fb").pack(side="left")

        f4 = tk.Frame(root, bg="#f2f6fb")
        f4.pack(fill="x", padx=16, pady=(6, 10))
        self.btn = tk.Button(f4, text="시작", width=14, bg="#00519E", fg="#fff",
                             font=("맑은 고딕", 10, "bold"), command=self.start)
        self.btn.pack(side="left")
        self.nbtn = tk.Button(f4, text="저장했어요 · 다음 ▶", width=18, state="disabled",
                              command=lambda: self.next_evt.set())
        self.nbtn.pack(side="left", padx=8)
        self.sbtn = tk.Button(f4, text="중단", width=8, state="disabled", command=self.do_stop)
        self.sbtn.pack(side="left")
        self.stat = tk.Label(f4, text="", bg="#f2f6fb", fg="#4a5a6b")
        self.stat.pack(side="left", padx=12)

        self.log = tk.Text(root, height=18, bg="#fff", relief="flat",
                           font=("맑은 고딕", 9), wrap="word")
        self.log.pack(fill="both", expand=True, padx=16, pady=(0, 16))
        self.say("엑셀 파일을 고르고 [시작]을 누르세요.")
        self.say("브라우저가 열리면 거기서 직접 로그인해 주세요. 로그인되면 자동으로 이어집니다.")

    def say(self, t):
        self.log.insert("end", t + "\n")
        self.log.see("end")
        self.root.update_idletasks()

    def pick(self):
        p = filedialog.askopenfilename(filetypes=[("엑셀", "*.xlsx")])
        if not p:
            return
        self.xl.delete(0, "end")
        self.xl.insert(0, p)
        try:
            self.items = read_excel(p)
            self.say("불러왔습니다 — 상품 %d건" % len(self.items))
        except Exception as e:
            messagebox.showerror(APP, "엑셀을 읽지 못했습니다.\n\n%s" % e)

    def make_tpl(self):
        p = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                         initialfile="고도몰_상품등록_양식.xlsx",
                                         filetypes=[("엑셀", "*.xlsx")])
        if not p:
            return
        try:
            write_template(p)
            self.say("양식을 만들었습니다 — %s" % p)
            messagebox.showinfo(APP, "양식 엑셀을 만들었습니다.\n내용을 채운 뒤 [열기]로 불러오세요.")
        except Exception as e:
            messagebox.showerror(APP, str(e))

    def do_stop(self):
        self.stop = True
        self.next_evt.set()
        self.say("중단 요청됨…")

    def start(self):
        if not self.items:
            messagebox.showwarning(APP, "엑셀 파일을 먼저 불러오세요.")
            return
        self.cfg["url"] = self.url.get().strip()
        save_cfg(self.cfg)
        self.stop = False
        self.btn.config(state="disabled")
        self.sbtn.config(state="normal")
        threading.Thread(target=self.run, daemon=True).start()

    def run(self):
        try:
            self.work()
        except Exception:
            self.say("오류가 났습니다:\n" + traceback.format_exc())
        finally:
            self.btn.config(state="normal")
            self.nbtn.config(state="disabled")
            self.sbtn.config(state="disabled")

    def work(self):
        from playwright.sync_api import sync_playwright
        base = self.url.get().strip().rstrip("/")
        reg = base + "/goods/goods_register.php"
        slow = 220 if self.slow.get() else 0
        with sync_playwright() as p:
            try:
                br = p.chromium.launch(headless=False, channel="chrome", slow_mo=slow,
                                       args=["--start-maximized"])
            except Exception:
                self.say("설치된 크롬을 못 찾아 기본 브라우저로 엽니다.")
                br = p.chromium.launch(headless=False, slow_mo=slow, args=["--start-maximized"])
            ctx = br.new_context(no_viewport=True)
            pg = ctx.new_page()
            self.say("브라우저를 엽니다…")
            pg.goto(base, wait_until="domcontentloaded")

            self.say("▶ 브라우저 창에서 직접 로그인해 주세요. (최대 10분 기다립니다)")
            ok = False
            for _ in range(600):
                if self.stop:
                    break
                try:
                    if pg.evaluate(LOGGED_IN_JS):
                        ok = True
                        break
                except Exception:
                    pass
                pg.wait_for_timeout(1000)
            if not ok:
                self.say("로그인을 확인하지 못했습니다. 중단합니다.")
                br.close()
                return
            self.say("로그인 확인됐습니다. 등록을 시작합니다.\n")

            total = len(self.items)
            for i, it in enumerate(self.items):
                if self.stop:
                    break
                self.stat.config(text="%d / %d" % (i + 1, total))
                self.say("[%d/%d] %s" % (i + 1, total, it.get("name", "")))
                pg.goto(reg, wait_until="domcontentloaded")
                pg.wait_for_timeout(1200)
                r = pg.evaluate(FILL_JS, it)
                if it.get("useStock", True) and it.get("stock") is not None:
                    pg.wait_for_timeout(300)
                    try:
                        pg.fill("input[name='stockCnt']", str(it["stock"]))
                    except Exception:
                        pass
                self.say("   채움 %d개%s" % (len(r.get("ok", [])),
                         ("  ·  못 채움: " + ", ".join(r.get("ng", []))) if r.get("ng") else ""))
                self.say("   → 이미지·옵션을 넣고 [저장]을 누른 뒤, 아래 [저장했어요 · 다음] 버튼을 눌러주세요.")
                self.nbtn.config(state="normal")
                self.next_evt.clear()
                self.next_evt.wait()
                self.nbtn.config(state="disabled")
                if self.stop:
                    break
            self.say("\n끝났습니다. 브라우저는 그대로 둡니다.")
            self.stat.config(text="완료")


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
